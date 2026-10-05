from io import BytesIO
import re
from typing import Any, cast

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from fastapi import (
    APIRouter,
    HTTPException,
)
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from services.draft_service import (
    DraftServiceError,
    get_draft_type,
    validate_draft_details,
)
from services.draft_legal_service import (
    DraftLegalServiceError,
    retrieve_draft_legal_requirements,
)
from services.draft_ai_service import (
    DraftAIServiceError,
    generate_draft,
)
from services.draft_validation_service import (
    DraftValidationError,
    validate_generated_draft,
)


router = APIRouter(
    prefix="/api/v1/draft",
    tags=["Draft"],
)


class DraftCreateRequest(BaseModel):
    draft_type: str = Field(
        ...,
        min_length=1,
    )

    details: dict[str, Any] = Field(
        default_factory=dict,
    )


class DraftDownloadRequest(BaseModel):
    title: str = Field(
        default="Legal Draft",
    )

    draft: str = Field(
        ...,
        min_length=1,
    )

    draft_type_name: str = Field(
        default="Legal Draft",
    )

    verified_sources: Any = Field(
        default_factory=list,
    )


def safe_filename(
    value: str,
) -> str:
    """
    Create a safe filename for the downloaded
    Word document.
    """

    cleaned = re.sub(
        r'[<>:"/\\|?*]',
        "",
        value,
    )

    cleaned = re.sub(
        r"\s+",
        "_",
        cleaned.strip(),
    )

    if not cleaned:
        return "legal_draft"

    return cleaned[:80]


def add_draft_text_to_document(
    document: DocxDocument,
    draft_text: str,
) -> None:
    """
    Add generated draft text to the Word document.
    """

    lines = draft_text.splitlines()

    for raw_line in lines:
        line = raw_line.strip()

        if not line:
            document.add_paragraph()
            continue

        is_numbered_heading = bool(
            re.match(
                r"^\d+[\.\)]\s+.+",
                line,
            )
        )

        is_upper_heading = (
            len(line) <= 100
            and line.upper() == line
            and any(
                character.isalpha()
                for character in line
            )
        )

        if (
            is_numbered_heading
            or is_upper_heading
        ):
            paragraph = (
                document.add_paragraph()
            )

            run = paragraph.add_run(
                line
            )

            run.bold = True
            run.font.size = Pt(12)

            continue

        paragraph = (
            document.add_paragraph(
                line
            )
        )

        paragraph.paragraph_format.space_after = (
            Pt(6)
        )


def get_verified_sources(
    value: Any,
) -> list[dict[str, Any]]:
    """
    Safely convert verified_sources into a typed
    list of dictionaries.
    """

    if not isinstance(
        value,
        list,
    ):
        return []

    raw_sources = cast(
        list[Any],
        value,
    )

    verified_sources: list[
        dict[str, Any]
    ] = []

    for item in raw_sources:
        if isinstance(
            item,
            dict,
        ):
            verified_sources.append(
                cast(
                    dict[str, Any],
                    item,
                )
            )

    return verified_sources


def create_word_document(
    request: DraftDownloadRequest,
) -> BytesIO:
    """
    Build the generated legal draft as a
    professional Microsoft Word document.
    """

    document: DocxDocument = Document()

    # ---------------------------------------------------------
    # Main title
    # ---------------------------------------------------------

    title = (
        request.title.strip()
        or request.draft_type_name.strip()
        or "Legal Draft"
    )

    title_paragraph = (
        document.add_paragraph()
    )

    title_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = (
        title_paragraph.add_run(
            title
        )
    )

    title_run.bold = True
    title_run.font.size = Pt(18)

    # ---------------------------------------------------------
    # Draft type subtitle
    # ---------------------------------------------------------

    type_name = (
        request.draft_type_name.strip()
    )

    if (
        type_name
        and type_name.lower()
        != title.lower()
    ):
        subtitle = (
            document.add_paragraph()
        )

        subtitle.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        subtitle_run = (
            subtitle.add_run(
                type_name
            )
        )

        subtitle_run.italic = True
        subtitle_run.font.size = Pt(10)

    document.add_paragraph()

    # ---------------------------------------------------------
    # Draft content
    # ---------------------------------------------------------

    add_draft_text_to_document(
        document=document,
        draft_text=request.draft,
    )

    # ---------------------------------------------------------
    # Verified Bahrain legal sources
    # ---------------------------------------------------------

    verified_sources = (
        get_verified_sources(
            request.verified_sources
        )
    )

    if verified_sources:
        document.add_page_break()

        sources_heading = (
            document.add_paragraph()
        )

        sources_run = (
            sources_heading.add_run(
                "Verified Bahrain Legal Sources"
            )
        )

        sources_run.bold = True
        sources_run.font.size = Pt(14)

        for index, source in enumerate(
            verified_sources,
            start=1,
        ):
            law_number = str(
                source.get(
                    "law_number",
                    "",
                )
            ).strip()

            year = str(
                source.get(
                    "year",
                    "",
                )
            ).strip()

            article = str(
                source.get(
                    "article",
                    "",
                )
            ).strip()

            source_url = str(
                source.get(
                    "source_url",
                    "",
                )
            ).strip()

            source_title = str(
                source.get(
                    "title",
                    "",
                )
            ).strip()

            parts: list[str] = []

            if source_title:
                parts.append(
                    source_title
                )

            if law_number:
                law_label = (
                    f"Law {law_number}"
                )

                if year:
                    law_label += (
                        f" / {year}"
                    )

                parts.append(
                    law_label
                )

            if article:
                parts.append(
                    f"Article {article}"
                )

            label = " — ".join(
                parts
            )

            if not label:
                label = (
                    f"Legal Source {index}"
                )

            paragraph = (
                document.add_paragraph(
                    style="List Number"
                )
            )

            label_run = (
                paragraph.add_run(
                    label
                )
            )

            label_run.bold = True

            if source_url:
                url_paragraph = (
                    document.add_paragraph(
                        source_url
                    )
                )

                url_paragraph.paragraph_format.left_indent = (
                    Pt(18)
                )

    # ---------------------------------------------------------
    # Footer note
    # ---------------------------------------------------------

    document.add_paragraph()

    footer_note = (
        document.add_paragraph()
    )

    footer_note.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = (
        footer_note.add_run(
            "Generated by Agentic AI"
        )
    )

    footer_run.italic = True
    footer_run.font.size = Pt(8)

    # ---------------------------------------------------------
    # Save Word document to memory
    # ---------------------------------------------------------

    buffer = BytesIO()

    document.save(
        buffer
    )

    buffer.seek(0)

    return buffer


@router.post("/create")
async def create_draft(
    request: DraftCreateRequest,
) -> dict[str, Any]:

    try:
        draft_type = (
            request.draft_type
            .strip()
            .lower()
        )

        draft_config = get_draft_type(
            draft_type
        )

        details = validate_draft_details(
            draft_type=draft_type,
            details=request.details,
        )

        legal_requirements = (
            retrieve_draft_legal_requirements(
                draft_type=draft_type,
                limit_per_query=3,
            )
        )

        ai_result = generate_draft(
            draft_type=draft_type,
            details=details,
            legal_requirements=(
                legal_requirements
            ),
        )

        validated_result = (
            validate_generated_draft(
                ai_result=ai_result,
                legal_requirements=(
                    legal_requirements
                ),
            )
        )

        return {
            "success": True,
            "message": (
                "Draft created successfully."
            ),
            "draft_type": draft_type,
            "draft_type_name": (
                draft_config["name"]
            ),
            "legal_requirements_retrieved": (
                len(legal_requirements)
            ),
            "result": validated_result,
        }

    except DraftServiceError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc),
        ) from exc

    except DraftLegalServiceError as exc:
        raise HTTPException(
            status_code=500,
            detail=str(exc),
        ) from exc

    except DraftAIServiceError as exc:
        raise HTTPException(
            status_code=502,
            detail=str(exc),
        ) from exc

    except DraftValidationError as exc:
        raise HTTPException(
            status_code=500,
            detail=str(exc),
        ) from exc

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=(
                "Draft creation failed: "
                f"{str(exc)}"
            ),
        ) from exc


@router.post("/download/word")
async def download_draft_word(
    request: DraftDownloadRequest,
) -> StreamingResponse:
    """
    Download a generated legal draft
    as a Microsoft Word document.
    """

    try:
        word_buffer = (
            create_word_document(
                request
            )
        )

        filename = (
            safe_filename(
                request.title
                or request.draft_type_name
            )
            + ".docx"
        )

        return StreamingResponse(
            word_buffer,
            media_type=(
                "application/"
                "vnd.openxmlformats-"
                "officedocument."
                "wordprocessingml.document"
            ),
            headers={
                "Content-Disposition": (
                    'attachment; filename="'
                    + filename
                    + '"'
                )
            },
        )

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=(
                "Word document creation "
                f"failed: {str(exc)}"
            ),
        ) from exc