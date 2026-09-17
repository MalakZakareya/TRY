from pathlib import Path
from io import BytesIO
from typing import Any, TypedDict

import pymupdf
from docx import Document


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_SIZE = 10 * 1024 * 1024


class DocumentProcessingError(Exception):
    pass


class DocumentResult(TypedDict):
    filename: str
    file_type: str
    file_size: int
    characters: int
    text: str


def validate_document(filename: str, content: bytes) -> str:
    if not filename:
        raise DocumentProcessingError(
            "File name is missing."
        )

    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise DocumentProcessingError(
            "Unsupported file type. Please upload PDF, DOCX, or TXT."
        )

    if not content:
        raise DocumentProcessingError(
            "The uploaded file is empty."
        )

    if len(content) > MAX_FILE_SIZE:
        raise DocumentProcessingError(
            "File is too large. Maximum size is 10 MB."
        )

    return extension


def extract_pdf_text(content: bytes) -> str:
    try:
        pdf_document: Any = pymupdf.open(
            stream=content,
            filetype="pdf"
        )

        pages: list[str] = []

        for page in pdf_document:
            raw_text: Any = page.get_text()

            if not isinstance(raw_text, str):
                continue

            cleaned_text = raw_text.strip()

            if cleaned_text:
                pages.append(cleaned_text)

        pdf_document.close()

        return "\n\n".join(pages)

    except Exception as exc:
        raise DocumentProcessingError(
            "Unable to read the PDF file."
        ) from exc


def extract_docx_text(content: bytes) -> str:
    try:
        document = Document(
            BytesIO(content)
        )

        paragraphs: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        for table in document.tables:
            for row in table.rows:
                row_text: list[str] = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:
                    paragraphs.append(
                        " | ".join(row_text)
                    )

        return "\n\n".join(paragraphs)

    except Exception as exc:
        raise DocumentProcessingError(
            "Unable to read the DOCX file."
        ) from exc


def extract_txt_text(content: bytes) -> str:
    try:
        return content.decode("utf-8")

    except UnicodeDecodeError:
        try:
            return content.decode(
                "utf-8-sig"
            )

        except UnicodeDecodeError as exc:
            raise DocumentProcessingError(
                "Unable to read the TXT file."
            ) from exc


def extract_document_text(
    filename: str,
    content: bytes
) -> DocumentResult:

    extension = validate_document(
        filename,
        content
    )

    if extension == ".pdf":
        text = extract_pdf_text(
            content
        )

    elif extension == ".docx":
        text = extract_docx_text(
            content
        )

    elif extension == ".txt":
        text = extract_txt_text(
            content
        )

    else:
        raise DocumentProcessingError(
            "Unsupported document type."
        )

    cleaned_text = text.strip()

    if not cleaned_text:
        raise DocumentProcessingError(
            "No readable text was found in this document."
        )

    return {
        "filename": filename,
        "file_type": extension.removeprefix("."),
        "file_size": len(content),
        "characters": len(cleaned_text),
        "text": cleaned_text,
    }